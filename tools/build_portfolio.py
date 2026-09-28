"""Build the portfolio: fill the course pro-forma, append the report, export a PDF via Word.

Usage: python tools/build_portfolio.py
Inputs:  docs/portfolio/proforma.docx, cover.json, report.md (+ figures referenced there)
Outputs: dist/portfolio/KG_Portfolio_Lindner-structured.{docx,pdf} and page previews (PNG)

Page numbers on the cover pages are found by exporting once, locating each section heading in
the PDF, and building a second time.
"""

import json
import re
import subprocess
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parents[1] / "docs" / "portfolio"
OUT = Path(__file__).resolve().parents[1] / "dist" / "portfolio"
NAME = "KG_Portfolio_Lindner-structured"

PRIMARY = RGBColor(0x0A, 0x4F, 0x47)
MUTED = RGBColor(0x5D, 0x6B, 0x64)
W14 = "http://schemas.microsoft.com/office/word/2010/wordml"
INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
IMAGE = re.compile(r"!\[(?P<caption>[^\]]*)\]\((?P<path>[^)]+)\)(\{width=(?P<width>[\d.]+)\})?")


# ---------------------------------------------------------------- styles


def _paragraph_style(
    doc,
    name,
    size,
    *,
    bold=False,
    italic=False,
    colour=None,
    font=None,
    before=0,
    after=4,
    outline=None,
    keep_next=False,
):
    style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = doc.styles["Normal"]
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if colour:
        style.font.color.rgb = colour
    if font:
        style.font.name = font
        style.element.rPr.rFonts.set(qn("w:eastAsia"), font)
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.keep_with_next = keep_next
    if outline is not None:
        outline_element = OxmlElement("w:outlineLvl")
        outline_element.set(qn("w:val"), str(outline))
        style.element.get_or_add_pPr().append(outline_element)
    return style


def add_styles(doc) -> None:
    body = _paragraph_style(doc, "Report Body", 10.5, after=5)
    body.paragraph_format.line_spacing = 1.08
    _paragraph_style(
        doc,
        "Heading 1",
        14,
        bold=True,
        colour=PRIMARY,
        before=12,
        after=4,
        outline=0,
        keep_next=True,
    )
    _paragraph_style(
        doc,
        "Heading 2",
        11.5,
        bold=True,
        colour=PRIMARY,
        before=8,
        after=3,
        outline=1,
        keep_next=True,
    )
    _paragraph_style(doc, "Report Code", 7.8, font="Consolas", after=0)
    _paragraph_style(doc, "Report Caption", 8.5, italic=True, colour=MUTED, after=8)
    _paragraph_style(doc, "Report Table", 8.5, after=0)
    _paragraph_style(doc, "Report Reference", 9, after=3)


# ---------------------------------------------------------------- cover pages


def set_checkbox(sdt, checked: bool) -> None:
    box = sdt.find(f".//{{{W14}}}checked")
    box.set(f"{{{W14}}}val", "1" if checked else "0")
    sdt.find(".//" + qn("w:t")).text = "☒" if checked else "☐"


def checkboxes(element) -> list:
    return [
        sdt for sdt in element.iter(qn("w:sdt")) if sdt.find(f".//{{{W14}}}checkbox") is not None
    ]


def replace_text(paragraph, text: str, *, highlight: bool = False) -> None:
    runs = paragraph.runs
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""
    if highlight:
        runs[0].font.highlight_color = WD_COLOR_INDEX.YELLOW


def replace_in_paragraph(paragraph, old: str, new: str) -> None:
    text = paragraph.text.replace(old, new)
    replace_text(paragraph, text)


def cell_text(cell, text: str, *, highlight: bool = False) -> None:
    paragraphs = [p for p in cell.paragraphs if p.runs]
    replace_text(paragraphs[0], text, highlight=highlight)
    for extra in paragraphs[1:]:
        extra._p.getparent().remove(extra._p)


def fill_cover(doc, cover: dict, pages: dict[str, int]) -> None:
    for paragraph in doc.paragraphs[:6]:
        if "<Your mini-project title>" in paragraph.text:
            replace_in_paragraph(paragraph, "<Your mini-project title>", cover["title"])
        if "<Your name>" in paragraph.text:
            replace_in_paragraph(paragraph, "<Your name>", cover["name"])
        if paragraph.text.startswith("Mode"):
            boxes = checkboxes(paragraph._p)
            set_checkbox(boxes[0], cover["mode"] == "6 ECTS")
            set_checkbox(boxes[1], cover["mode"] == "3 ECTS")

    for index, table in enumerate(doc.tables[:12], start=1):
        entry = cover["learning_outcomes"][f"LO{index}"]
        basic, exceeded = checkboxes(table.rows[0].cells[1]._tc)
        set_checkbox(basic, entry["level"] == "basic")
        set_checkbox(exceeded, entry["level"] == "exceeded")
        cell_text(table.rows[1].cells[0], entry["text"])
        cell_text(table.rows[1].cells[1], _page_reference(entry["sections"], pages))

    hours = [cover["hours_project"], cover["hours_document"]]
    for table, value in zip(doc.tables[12:14], hours, strict=True):
        cell = table.rows[0].cells[1]
        cell_text(cell, f"{value} hours" if value else "<XX*> hours", highlight=value is None)

    set_checkbox(checkboxes(doc.tables[14]._tbl)[0], cover["declaration_confirmed"])
    for table, key, noun in (
        (doc.tables[15], "ai_project", "of the mini-project"),
        (doc.tables[16], "ai_document", "of this document"),
    ):
        info = cover[key]
        set_checkbox(checkboxes(table.rows[0].cells[1]._tc)[0], info["used"])
        percent = info["percent"]
        percent_paragraph = [p for p in table.rows[0].cells[1].paragraphs if "%" in p.text][0]
        amount = f"{percent}%" if percent is not None else "<XX>%"
        lead = " I used it for roughly " if "I used it" in percent_paragraph.text else ""
        replace_text(percent_paragraph, f"{lead}{amount} {noun}", highlight=percent is None)
        cell_text(table.rows[1].cells[0], info["text"], highlight="[" in info["text"])


def _page_reference(sections: list[str], pages: dict[str, int]) -> str:
    if not sections:
        return "–"
    parts = []
    for section in sections:
        label = section if section.startswith("Appendix") else f"Section {section}"
        page = pages.get(section)
        parts.append(f"{label} (p. {page})" if page else label)
    return ", ".join(parts)


def remove_template_guidance(doc) -> None:
    """Drop everything after the last cover table (breaks, empty lines, writing guidance)."""
    body = doc.element.body
    last_table = doc.tables[-1]._tbl
    for element in list(body)[list(body).index(last_table) + 1 :]:
        if not element.tag.endswith("}sectPr"):
            body.remove(element)


# ---------------------------------------------------------------- report


def add_inline(paragraph, text: str) -> None:
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif token.startswith("*"):
            paragraph.add_run(token[1:-1]).italic = True
        else:
            paragraph.add_run(token)


def shade(paragraph, colour: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), colour)
    properties.append(fill)


def add_table(doc, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = doc.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    lengths = [max(len(row[i]) for row in rows) for i in range(len(rows[0]))]
    total = sum(min(length, 60) + 8 for length in lengths)
    usable = 16.0
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        row_properties.append(OxmlElement("w:cantSplit"))
    for i, row in enumerate(rows):
        for j, text in enumerate(row):
            cell = table.cell(i, j)
            cell.width = Cm(usable * (min(lengths[j], 60) + 8) / total)
            paragraph = cell.paragraphs[0]
            paragraph.style = doc.styles["Report Table"]
            add_inline(paragraph, text)
            if i == 0:
                for run in paragraph.runs:
                    run.bold = True
                cell_properties = cell._tc.get_or_add_tcPr()
                fill = OxmlElement("w:shd")
                fill.set(qn("w:val"), "clear")
                fill.set(qn("w:color"), "auto")
                fill.set(qn("w:fill"), "D8EDE9")
                cell_properties.append(fill)
    doc.add_paragraph(style="Report Body").paragraph_format.space_after = Pt(2)


def add_report(doc, markdown: str) -> None:
    lines = markdown.splitlines()
    i = 0
    first_heading = True
    while i < len(lines):
        line = lines[i]
        if line.startswith("# "):
            heading = doc.add_paragraph(line[2:], style="Heading 1")
            heading.paragraph_format.page_break_before = first_heading or line.startswith(
                "# Appendix"
            )
            first_heading = False
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 2")
        elif line.startswith("```"):
            i += 1
            block = []
            while not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            for code_line in block:
                paragraph = doc.add_paragraph(code_line or " ", style="Report Code")
                paragraph.paragraph_format.left_indent = Cm(0.3)
                shade(paragraph, "F1F3F1")
            doc.add_paragraph(style="Report Code").paragraph_format.space_after = Pt(3)
        elif line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        elif match := IMAGE.fullmatch(line.strip()):
            path = (HERE / match["path"]).resolve()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            paragraph.add_run().add_picture(str(path), width=Cm(float(match["width"] or 15)))
            caption = doc.add_paragraph(style="Report Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(caption, match["caption"])
        elif re.match(r"Table \d+:", line):
            caption = doc.add_paragraph(style="Report Caption")
            caption.paragraph_format.keep_with_next = True
            caption.paragraph_format.space_after = Pt(3)
            add_inline(caption, line)
        elif re.match(r"\d+\. ", line):
            paragraph = doc.add_paragraph(style="Report Body")
            paragraph.paragraph_format.left_indent = Cm(0.6)
            paragraph.paragraph_format.first_line_indent = Cm(-0.6)
            add_inline(paragraph, line)
        elif line.startswith("["):
            add_inline(doc.add_paragraph(style="Report Reference"), line)
        elif line.strip():
            add_inline(doc.add_paragraph(style="Report Body"), line)
        i += 1


def add_page_numbers(doc) -> None:
    footer = doc.sections[0].footer
    paragraph = footer.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)
    run.font.size = Pt(9)


# ---------------------------------------------------------------- build


def build(pages: dict[str, int]) -> Path:
    cover = json.loads((HERE / "cover.json").read_text(encoding="utf-8"))
    doc = Document(str(HERE / "proforma.docx"))
    add_styles(doc)
    fill_cover(doc, cover, pages)
    remove_template_guidance(doc)
    add_report(doc, (HERE / "report.md").read_text(encoding="utf-8"))
    add_page_numbers(doc)
    OUT.mkdir(parents=True, exist_ok=True)
    target = OUT / f"{NAME}.docx"
    doc.save(target)
    return target


def export_pdf(docx_path: Path) -> Path:
    pdf_path = docx_path.with_suffix(".pdf")
    script = (
        "$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
        f"$d = $w.Documents.Open('{docx_path}'); "
        "$d.Fields.Update() | Out-Null; "
        f"$d.SaveAs2('{pdf_path}', 17); $d.Close($false); $w.Quit()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True)
    return pdf_path


def section_pages(pdf_path: Path) -> dict[str, int]:
    """Page number of every numbered heading and the appendix, searched from the report's start."""
    keys = []
    for line in (HERE / "report.md").read_text(encoding="utf-8").splitlines():
        if match := re.match(r"#{1,2} (\d+(?:\.\d+)?|Appendix A)\b", line):
            keys.append(match.group(1))
    pages: dict[str, int] = {}
    with fitz.open(pdf_path) as pdf:
        texts = [[text.strip() for text in page.get_text().splitlines()] for page in pdf]
    start = next(i for i, page in enumerate(texts) if "1 Scenario" in page)
    for number, page in enumerate(texts[start:], start=start + 1):
        for text in page:
            for key in keys:
                if key not in pages and text.startswith((f"{key} ", f"{key}:")):
                    pages[key] = number
    return pages


def previews(pdf_path: Path) -> list[Path]:
    paths = []
    with fitz.open(pdf_path) as pdf:
        for number, page in enumerate(pdf, start=1):
            target = OUT / f"page-{number:02d}.png"
            page.get_pixmap(dpi=80).save(target)
            paths.append(target)
    return paths


def main() -> None:
    first = export_pdf(build({}))
    pages = section_pages(first)
    final = export_pdf(build(pages))
    with fitz.open(final) as pdf:
        count = pdf.page_count
    previews(final)
    print(f"{final} ({count} pages); section pages: {pages}")


if __name__ == "__main__":
    main()
